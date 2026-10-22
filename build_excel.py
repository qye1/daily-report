# Excel отчет: данные по дням лежат листом, а сводка, изменения и доли считаются живыми формулами,
# чтобы человек мог менять дату в ячейке B2 и видеть пересчет

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

F = "Arial"
HEAD = PatternFill("solid", fgColor="2B6CB0")
INPUT = PatternFill("solid", fgColor="FFF5C2")
GREEN = PatternFill("solid", fgColor="C6F6D5")
AMBER = PatternFill("solid", fgColor="FEEBC8")
RED = PatternFill("solid", fgColor="FED7D7")
thin = Side(style="thin", color="CBD5E0")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
STATUS_FILL = {"ok": GREEN, "warn": AMBER, "fail": RED}
STATUS_RU = {"ok": "в порядке", "warn": "внимание", "fail": "ошибка"}
MONEY = '#,##0'
PCT = '0.0%;[Red]-0.0%;0.0%'


def _head(ws, row, titles):
    for j, t in enumerate(titles, start=1):
        c = ws.cell(row=row, column=j, value=t)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX


def _style(ws, rows, cols, fmt=None):
    for r in rows:
        for j in cols:
            c = ws.cell(row=r, column=j)
            c.font = Font(name=F, size=10)
            c.border = BOX
            if fmt and j in fmt:
                c.number_format = fmt[j]


def build_workbook(path, daily, regions, categories, checks, lines, anomaly, rd, cfg):
    wb = Workbook()
    n = cfg["history_days_in_excel"]
    hist = daily.tail(n)

    # лист с данными по дням
    wd = wb.active
    wd.title = "Дни"
    _head(wd, 1, ["Дата", "Выручка", "Заказы", "Покупатели", "Средний чек", "Выручка, среднее за 7 дней", "К тому же дню неделей раньше"])
    for i, (d, r) in enumerate(hist.iterrows(), start=2):
        wd.cell(row=i, column=1, value=d.to_pydatetime()).number_format = "DD.MM.YYYY"
        wd.cell(row=i, column=2, value=round(float(r["revenue"]), 2))
        wd.cell(row=i, column=3, value=int(r["orders"]))
        wd.cell(row=i, column=4, value=int(r["customers"]))
        wd.cell(row=i, column=5, value=f"=IF(C{i}=0,0,B{i}/C{i})")
        if i >= 8:
            wd.cell(row=i, column=6, value=f"=AVERAGE(B{i - 6}:B{i})")
        if i >= 9:
            wd.cell(row=i, column=7, value=f"=IF(B{i - 7}=0,\"\",B{i}/B{i - 7}-1)")
    last = len(hist) + 1
    _style(wd, range(2, last + 1), range(1, 8), {1: "DD.MM.YYYY", 2: MONEY, 3: "#,##0", 4: "#,##0", 5: MONEY, 6: MONEY, 7: PCT})
    wd.conditional_formatting.add(f"G9:G{last}", CellIsRule(operator="lessThan", formula=["-0.2"], fill=RED))
    wd.conditional_formatting.add(f"G9:G{last}", CellIsRule(operator="greaterThan", formula=["0.2"], fill=GREEN))
    for col, w in zip("ABCDEFG", [12, 14, 10, 12, 13, 18, 18]):
        wd.column_dimensions[col].width = w
    wd.row_dimensions[1].height = 42
    wd.freeze_panes = "B2"

    # сводка
    ws = wb.create_sheet("Сводка", 0)
    ws["A1"] = "Ежедневный отчет по продажам"
    ws["A1"].font = Font(name=F, size=16, bold=True, color="2B6CB0")
    ws["A2"] = "Отчетная дата (можно менять)"
    ws["A2"].font = Font(name=F, size=10)
    ws["B2"] = rd.to_pydatetime()
    ws["B2"].number_format = "DD.MM.YYYY"
    ws["B2"].fill = INPUT
    ws["B2"].font = Font(name=F, size=10, bold=True, color="0000FF")
    ws["B2"].border = BOX
    _head(ws, 4, ["Показатель", "Отчетный день", "Неделей раньше", "Изменение"])
    A, B, C, D = "Дни!$A:$A", "Дни!$B:$B", "Дни!$C:$C", "Дни!$E:$E"
    look = lambda col, off: f"=INDEX({col},MATCH($B$2-{off},{A},0))" if off else f"=INDEX({col},MATCH($B$2,{A},0))"
    rows = [("Выручка", B, MONEY), ("Заказы", C, "#,##0"), ("Средний чек", D, MONEY), ("Покупатели", "Дни!$D:$D", "#,##0")]
    for i, (name, col, fmt) in enumerate(rows, start=5):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=look(col, 0)).number_format = fmt
        ws.cell(row=i, column=3, value=look(col, 7)).number_format = fmt
        ws.cell(row=i, column=4, value=f"=IF(C{i}=0,\"\",B{i}/C{i}-1)").number_format = PCT
    _style(ws, range(5, 9), range(1, 5), {2: MONEY, 3: MONEY, 4: PCT})
    for i, fmt in zip(range(5, 9), [MONEY, "#,##0", MONEY, "#,##0"]):
        ws.cell(row=i, column=2).number_format = fmt
        ws.cell(row=i, column=3).number_format = fmt
    _head(ws, 10, ["Период", "Выручка", "Предыдущий период", "Изменение"])
    ws["A11"], ws["A12"] = "Последние 7 дней", "С начала месяца"
    ws["B11"] = f'=SUMIFS({B},{A},">"&$B$2-7,{A},"<="&$B$2)'
    ws["C11"] = f'=SUMIFS({B},{A},">"&$B$2-14,{A},"<="&$B$2-7)'
    ws["B12"] = f'=SUMIFS({B},{A},">="&DATE(YEAR($B$2),MONTH($B$2),1),{A},"<="&$B$2)'
    ws["C12"] = f'=SUMIFS({B},{A},">="&EDATE(DATE(YEAR($B$2),MONTH($B$2),1),-1),{A},"<="&EDATE($B$2,-1))'
    for r in (11, 12):
        ws[f"D{r}"] = f'=IF(C{r}=0,"",B{r}/C{r}-1)'
    _style(ws, range(11, 13), range(1, 5), {2: MONEY, 3: MONEY, 4: PCT})
    ws["A13"] = "Для месяца нужна история с начала прошлого месяца: в таблице хранится последние " + str(n) + " дней."
    ws["A13"].font = Font(name=F, size=8, italic=True, color="718096")

    ws["A15"] = "Выводы"
    ws["A15"].font = Font(name=F, size=12, bold=True)
    for i, line in enumerate(lines, start=16):
        ws.cell(row=i, column=1, value=line).font = Font(name=F, size=10)
    r = 16 + len(lines) + 1
    ws.cell(row=r, column=1, value="Выводы записаны на момент сборки отчета для даты из названия файла, остальное на листе пересчитывается.").font = \
        Font(name=F, size=8, italic=True, color="718096")
    for col, w in zip("ABCD", [34, 18, 20, 14]):
        ws.column_dimensions[col].width = w

    ch = LineChart()
    ch.title = f"Выручка за {cfg['chart_days']} дней"
    ch.height, ch.width = 7.5, 17
    k = cfg["chart_days"]
    first = max(2, last - k + 1)
    ch.add_data(Reference(wd, min_col=2, min_row=first, max_row=last), titles_from_data=False)
    ch.add_data(Reference(wd, min_col=6, min_row=first, max_row=last), titles_from_data=False)
    ch.set_categories(Reference(wd, min_col=1, min_row=first, max_row=last))
    ch.series[0].tx = SeriesLabel(v="Выручка")
    ch.series[1].tx = SeriesLabel(v="Среднее за 7 дней")
    ch.x_axis.number_format = "DD.MM"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.number_format = "#,##0"
    ws.add_chart(ch, f"F4")

    # разрезы
    for title, t in [("Регионы", regions), ("Категории", categories)]:
        w = wb.create_sheet(title)
        _head(w, 1, [title[:-1] if title == "Категории" else "Регион", "Отчетный день", "Неделей раньше", "Изменение, сумма", "Изменение, %", "Доля в выручке"])
        if title == "Категории":
            w["A1"] = "Категория"
        end = len(t) + 1
        for i, (name, r_) in enumerate(t.iterrows(), start=2):
            w.cell(row=i, column=1, value=str(name))
            w.cell(row=i, column=2, value=round(float(r_["today"]), 2))
            w.cell(row=i, column=3, value=round(float(r_["week_ago"]), 2))
            w.cell(row=i, column=4, value=f"=B{i}-C{i}")
            w.cell(row=i, column=5, value=f'=IF(C{i}=0,"",B{i}/C{i}-1)')
            w.cell(row=i, column=6, value=f"=IF(SUM(B$2:B${end})=0,0,B{i}/SUM(B$2:B${end}))")
        tr = end + 1
        w.cell(row=tr, column=1, value="Итого")
        for j, col in zip((2, 3, 4), "BCD"):
            w.cell(row=tr, column=j, value=f"=SUM({col}2:{col}{end})")
        w.cell(row=tr, column=5, value=f'=IF(C{tr}=0,"",B{tr}/C{tr}-1)')
        w.cell(row=tr, column=6, value=f"=SUM(F2:F{end})")
        _style(w, range(2, tr + 1), range(1, 7), {2: MONEY, 3: MONEY, 4: MONEY, 5: PCT, 6: "0.0%"})
        for j in range(1, 7):
            w.cell(row=tr, column=j).font = Font(name=F, size=10, bold=True)
        w.conditional_formatting.add(f"E2:E{end}", CellIsRule(operator="lessThan", formula=["-0.2"], fill=RED))
        w.conditional_formatting.add(f"E2:E{end}", CellIsRule(operator="greaterThan", formula=["0.2"], fill=GREEN))
        for col, wd_ in zip("ABCDEF", [22, 16, 16, 16, 14, 14]):
            w.column_dimensions[col].width = wd_
        w.row_dimensions[1].height = 32

    # проверки
    wc = wb.create_sheet("Проверки данных")
    _head(wc, 1, ["Проверка", "Статус", "Подробности"])
    for i, c in enumerate(checks, start=2):
        wc.cell(row=i, column=1, value=c.name)
        s = wc.cell(row=i, column=2, value=STATUS_RU[c.status])
        s.fill = STATUS_FILL[c.status]
        wc.cell(row=i, column=3, value=c.detail)
    _style(wc, range(2, len(checks) + 2), range(1, 4))
    for col, w_ in zip("ABC", [30, 14, 90]):
        wc.column_dimensions[col].width = w_
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
